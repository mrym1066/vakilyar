import os
import uuid
#import re
#import tempfile
from collections import defaultdict
from datetime import datetime, date, time, timedelta
from decimal import Decimal
#from urllib.parse import quote
import urllib.parse

import jdatetime
#import jalali_date

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.db import transaction
from django.db.models import Count, Q, Sum
from django.http import FileResponse, HttpRequest, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
#from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from django.utils.dateparse import parse_datetime
from django.views.generic import DetailView, ListView
from .models import MovakelType  

# PDF/Word
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from .pdf_utils import render_to_pdf, get_common_pdf_context, safe_pdf_filename

from reportlab.pdfgen import canvas

try:
    import arabic_reshaper
    from bidi.algorithm import get_display
    BIDI_AVAILABLE = True
except ImportError:
    BIDI_AVAILABLE = False

# Local imports
from movakel_module.backup_utils import create_backup, get_all_backups, delete_backup
from site_module.models import SiteSetting
#from utils.http_service import get_client_ip
#from utils.convertors import group_list

from .forms import (
    DamageCalculationForm, RequestMeetingForm, ServiceTypeForm,
    RequestMeetingEditForm, VisitEditForm, DefenseDocumentEditForm,
    PDFFileEditForm, MovakelBasicEditForm, CourtInfoEditForm,
    MovakelNotesEditForm, MovakelForm,
)
from .models import (
    Movakel, MovakelCategory, InstallmentPayment, MovakelGallery,
    MovakelPayment, ServiceType, DefenseDocument, Branch, RequestMeeting,
    PaymentDetail, Visit, PDFFile,PAYMENT_FOR_CHOICES,
)
from .utils import calculate_damage


# ============================================================
# Helper
# ============================================================
def is_admin(user):
     #بررسی اینکه کاربر ادمین است یا نه
    return user.is_authenticated and user.is_superuser


# ============================================================
# List / Detail
# ============================================================
@login_required(login_url='/login/')
def movakel_list_view(request):
    """جایگزین MovakelListView برای امکان استفاده ساده‌تر"""
    return MovakelListView.as_view()(request)


class MovakelListView(LoginRequiredMixin, ListView):
    template_name = 'movakel_module/movakel_list.html'
    model = Movakel
    context_object_name = 'movakels'
    ordering = ['-contract_date']
    paginate_by = 6
    login_url = '/login/'

    def get_queryset(self):
        query = super().get_queryset().filter(is_delete=False)

        start_total_amount = self.request.GET.get('start_total_amount')
        end_total_amount = self.request.GET.get('end_total_amount')

        if start_total_amount:
            query = query.filter(total_amount__gte=start_total_amount)
        if end_total_amount:
            query = query.filter(total_amount__lte=end_total_amount)

        category_name = self.kwargs.get('cat')
        type_name = self.kwargs.get('type')

        #if type_name:
           # query = query.filter(type__url_title__iexact=type_name)
        if category_name:
            query = query.filter(category__url_title__iexact=category_name)

        return query

    def get_context_data(self, *, object_list=None, **kwargs):
        context = super().get_context_data(**kwargs)
        movakels = context['movakels']  # ← همان queryset فیلترشده
        context['total_movakels'] = Movakel.objects.filter(is_delete=False).count()
        # دسته‌بندی بر اساس سال شمسی
        categorized = defaultdict(list)
        for m in movakels:
            if m.contract_date:
                try:
                    year = jdatetime.date.fromgregorian(date=m.contract_date).year
                except Exception:
                    year = m.contract_date.year
            else:
                year = jdatetime.date.today().year
            categorized[year].append(m)

        context['categorized_movakels'] = dict(sorted(categorized.items(), reverse=True))

        # دسته‌بندی‌ها
        categories = MovakelCategory.objects.all()
        category_counts = (
            Movakel.objects.filter(is_delete=False)
            .values('category__title')
            .annotate(count=Count('id'))
        )
        counts = {item['category__title']: item['count'] for item in category_counts}
        context['category_list'] = [
            {'name': c.title, 'count': counts.get(c.title, 0)}
            for c in categories
        ]

    # نوع پرونده‌ها
        type_counts = (
            Movakel.objects.filter(is_delete=False, type__isnull=False)
            .values('type__id', 'type__title', 'type__url_title')
            .annotate(count=Count('id'))
            .order_by('-count')
        )

        context['type_list'] = [
            {
                'id': item['type__id'],
                'title': item['type__title'],
                'url_title': item['type__url_title'],
                'count': item['count'],
            }
            for item in type_counts
]
        context['start_total_amount'] = self.request.GET.get('start_total_amount') or 0
        context['end_total_amount'] = self.request.GET.get('end_total_amount') or 0

        return context


class MovakelDetailView(LoginRequiredMixin, UserPassesTestMixin, DetailView):
    template_name = 'movakel_module/movakel_detail.html'
    model = Movakel
    context_object_name = 'movakel'
    slug_field = 'slug'
    slug_url_kwarg = 'slug'
    login_url = '/login/'

    def test_func(self):
        return self.request.user.is_authenticated and self.request.user.is_superuser

    def handle_no_permission(self):
        if not self.request.user.is_authenticated:
            return super().handle_no_permission()
        return render(self.request, 'errors/access_denied.html', status=403)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        movakel = self.get_object()

        context['pdf_files'] = movakel.pdf_files.all()

        # پرداخت‌ها
        payments = movakel.movakelpayments.filter(is_delete=False)
        context['total_received_amount'] = (
            payments.aggregate(total=Sum('amount'))['total'] or 0
        )
        context['payments_by_category'] = payments.values('payment_for').annotate(total=Sum('amount'))
        context['payments'] = payments.order_by('-payment_date')

        # نوع خدمت
        try:
            service_type = movakel.service_type
        except ServiceType.DoesNotExist:
            service_type = None
        context['selected_service_type'] = service_type
        context['total_service_fee'] = service_type.total_fee if service_type else 0

        # لایحه‌ها
        context['defense_documents'] = DefenseDocument.objects.filter(movakel=movakel)

        # اطلاعات دادگاه
        for field in (
            'hearing_date', 'primary_court_number', 'primary_court_notification_date',
            'appeal_court_number', 'appeal_court_notification_date',
            'executive_case_number', 'executive_case_notification_date',
            'mechanized_number',
        ):
            context[field] = getattr(movakel, field, None)

        # ملاقات‌ها
        context['meetings'] = movakel.meetings.all()
        context['total_meeting_fee'] = (
            movakel.meetings.aggregate(total_fee=Sum('meeting_fee'))['total_fee'] or 0
        )

        # پرونده‌های مرتبط
        related = Movakel.objects.filter(is_delete=False).exclude(id=movakel.id)
        if movakel.national_id:
            related = related.filter(
                Q(national_id=movakel.national_id) | Q(name__icontains=movakel.name)
            )
        else:
            related = related.filter(name__icontains=movakel.name)
        related = related.order_by('-contract_date')
        context['related_movakels'] = [
            related[i:i + 3] for i in range(0, len(related), 3)
        ]

        # مراجعات
        context['visits'] = movakel.visits.all().order_by('visit_date')

        # اقساط
        context['installment_payments'] = movakel.installment_payments.all()

        return context


# ============================================================
# Meeting Views
# ============================================================
@login_required(login_url='/login/')
def meeting_list_view(request):
    meetings = RequestMeeting.objects.all().order_by('-meeting_date')

    # ===== فیلتر جستجو =====
    query = request.GET.get('q', '').strip()
    status = request.GET.get('status', '').strip()
    meeting_type = request.GET.get('type', '').strip()

    if query:
        meetings = meetings.filter(
            Q(full_name__icontains=query) |
            Q(national_id__icontains=query) |
            Q(mobile_number__icontains=query)
        )

    if status == 'pending':
        meetings = meetings.filter(status='pending')
    elif status == 'converted':
        meetings = meetings.filter(status='converted')

    if meeting_type:
        meetings = meetings.filter(meeting_type=meeting_type)

    # ===== پردازش تاریخ‌ها =====
    for m in meetings:
        if isinstance(m.meeting_date, str):
            try:
                m.meeting_date = parse_datetime(m.meeting_date)
            except (ValueError, TypeError):
                m.meeting_date = None

        if isinstance(m.meeting_date, datetime):
            try:
                jd = jdatetime.datetime.fromgregorian(datetime=m.meeting_date)
                m.meeting_date_shamsi = jd.strftime('%Y/%m/%d')
                m.formatted_time = jd.strftime('%H:%M')
            except Exception:
                m.meeting_date_shamsi = "-"
                m.formatted_time = "-"
        else:
            m.meeting_date_shamsi = "-"
            m.formatted_time = "-"

        m.meeting_fee = m.meeting_fee if m.meeting_fee is not None else 0
        m.subjects_display = ", ".join(s.name for s in m.meeting_subject.all())

    return render(request, 'movakel_module/meetings.html', {
        'meetings': meetings,
        'query': query,
        'status_filter': status,
        'type_filter': meeting_type,
    })



@login_required(login_url='/login/')
def request_meeting_view(request, meeting_id=None):
    meeting = get_object_or_404(RequestMeeting, id=meeting_id) if meeting_id else None

    if request.method == "POST":
        form = RequestMeetingForm(request.POST, instance=meeting)
        if form.is_valid():
            form.save()
            messages.success(request, "✅ قرار ملاقات با موفقیت ثبت شد.")
            return redirect("meeting_list")
    else:
        form = RequestMeetingForm(instance=meeting)

    shamsi_date = jdatetime.date.today().strftime("%Y/%m/%d")
    return render(request, "movakel_module/request_meeting_form.html", {
        "form": form,
        "meeting": meeting,
        "shamsi_date": shamsi_date,
    })


@login_required(login_url='/login/')
def create_request_meeting(request):
    if request.method == 'POST':
        form = RequestMeetingForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, '✅ درخواست ملاقات با موفقیت ثبت شد.')
            return redirect('meeting_list')
    else:
        form = RequestMeetingForm()

    return render(request, 'movakel_module/request_meeting_form.html', {'form': form})


# ============================================================
# Edit Views
# ============================================================
@login_required(login_url='/login/')
def edit_request_meeting(request, pk):
    meeting = get_object_or_404(RequestMeeting, pk=pk)
    if request.method == 'POST':
        form = RequestMeetingEditForm(request.POST, instance=meeting)
        if form.is_valid():
            form.save()
            messages.success(request, '✅ اطلاعات ملاقات با موفقیت ویرایش شد.')
            if meeting.movakel_id:
                return redirect('movakel-detail', slug=meeting.movakel.slug)
            return redirect('meeting_list')
    else:
        form = RequestMeetingEditForm(instance=meeting)

    return render(request, 'movakel_module/related_edit.html', {
        'form': form,
        'title': 'ویرایش اطلاعات ملاقات',
        'subtitle': meeting.full_name,
        'back_url': meeting.movakel.get_absolute_url() if meeting.movakel else None,
    })


@login_required(login_url='/login/')
def edit_visit(request, pk):
    visit = get_object_or_404(Visit, pk=pk)
    if request.method == 'POST':
        form = VisitEditForm(request.POST, instance=visit)
        if form.is_valid():
            form.save()
            messages.success(request, '✅ اطلاعات مراجعه با موفقیت ویرایش شد.')
            return redirect('movakel-detail', slug=visit.movakel.slug)
    else:
        form = VisitEditForm(instance=visit)

    return render(request, 'movakel_module/related_edit.html', {
        'form': form,
        'title': 'ویرایش مراجعه موکل به دفتر',
        'subtitle': visit.movakel.name,
        'back_url': visit.movakel.get_absolute_url(),
    })


@login_required(login_url='/login/')
def service_type_detail(request, pk):
    service_type = get_object_or_404(ServiceType, pk=pk)
    if request.method == 'POST':
        form = ServiceTypeForm(request.POST, instance=service_type)
        if form.is_valid():
            form.save()
            messages.success(request, '✅ اطلاعات خدمات با موفقیت ویرایش شد.')
            return redirect('movakel-detail', slug=service_type.movakel.slug)
    else:
        form = ServiceTypeForm(instance=service_type)

    return render(request, 'movakel_module/related_edit.html', {
        'form': form,
        'service_type': service_type,

    })


@login_required(login_url='/login/')
def edit_defense_document(request, pk):
    doc = get_object_or_404(DefenseDocument, pk=pk)
    if request.method == 'POST':
        form = DefenseDocumentEditForm(request.POST, request.FILES, instance=doc)
        if form.is_valid():
            form.save()
            messages.success(request, '✅ لایحه با موفقیت ویرایش شد.')
            return redirect('movakel-detail', slug=doc.movakel.slug)
    else:
        form = DefenseDocumentEditForm(instance=doc)

    return render(request, 'movakel_module/related_edit.html', {
        'form': form,
        'title': 'ویرایش لایحه',
        'subtitle': doc.movakel.name,
        'back_url': doc.movakel.get_absolute_url(),
    })


@login_required(login_url='/login/')
def edit_pdf_file(request, pk):
    pdf = get_object_or_404(PDFFile, pk=pk)
    if request.method == 'POST':
        form = PDFFileEditForm(request.POST, request.FILES, instance=pdf)
        if form.is_valid():
            form.save()
            messages.success(request, '✅ فایل PDF با موفقیت ویرایش شد.')
            movakel = pdf.movakels.filter(is_delete=False).first()
            if movakel:
                return redirect('movakel-detail', slug=movakel.slug)
            return redirect('movakel-list')
    else:
        form = PDFFileEditForm(instance=pdf)

    movakel = pdf.movakels.filter(is_delete=False).first()
    return render(request, 'movakel_module/related_edit.html', {
        'form': form,
        'title': 'ویرایش فایل PDF',
        'subtitle': pdf.name,
        'back_url': movakel.get_absolute_url() if movakel else None,
    })


@login_required(login_url='/login/')
def edit_movakel(request, slug):
    movakel = get_object_or_404(Movakel, slug=slug, is_delete=False)
    if request.method == 'POST':
        form = MovakelBasicEditForm(request.POST, request.FILES, instance=movakel)
        if form.is_valid():
            form.save()
            messages.success(request, '✅ اطلاعات اصلی موکل با موفقیت ویرایش شد.')
            return redirect('movakel-detail', slug=movakel.slug)
    else:
        form = MovakelBasicEditForm(instance=movakel)

    return render(request, 'movakel_module/related_edit.html', {
        'form': form,
        'title': 'ویرایش اطلاعات اصلی موکل',
        'subtitle': movakel.name,
        'back_url': movakel.get_absolute_url(),
    })


@login_required(login_url='/login/')
def edit_court_info(request, slug):
    movakel = get_object_or_404(Movakel, slug=slug, is_delete=False)
    if request.method == 'POST':
        form = CourtInfoEditForm(request.POST, instance=movakel)
        if form.is_valid():
            form.save()
            messages.success(request, '✅ اطلاعات پرونده و دادگاه با موفقیت ویرایش شد.')
            return redirect('movakel-detail', slug=movakel.slug)
    else:
        form = CourtInfoEditForm(instance=movakel)

    return render(request, 'movakel_module/related_edit.html', {
        'form': form,
        'title': 'ویرایش اطلاعات پرونده و دادگاه',
        'subtitle': movakel.name,
        'back_url': movakel.get_absolute_url(),
    })


@login_required(login_url='/login/')
def edit_movakel_notes(request, slug):
    movakel = get_object_or_404(Movakel, slug=slug, is_delete=False)
    if request.method == 'POST':
        form = MovakelNotesEditForm(request.POST, instance=movakel)
        if form.is_valid():
            form.save()
            messages.success(request, '✅ ملاحظات پرونده با موفقیت ویرایش شد.')
            return redirect('movakel-detail', slug=movakel.slug)
    else:
        form = MovakelNotesEditForm(instance=movakel)

    return render(request, 'movakel_module/related_edit.html', {
        'form': form,
        'title': 'ویرایش ملاحظات پرونده',
        'subtitle': movakel.name,
        'back_url': movakel.get_absolute_url(),
    })


# ============================================================
# Service / Payment
# ============================================================
@login_required(login_url='/login/')
def get_service_price(request, service_type_id):
    try:
        st = ServiceType.objects.get(id=service_type_id)
        price = getattr(st, 'price', None)
        if price is None:
            return JsonResponse({'error': 'Price field not found'}, status=400)
        return JsonResponse({'price': price})
    except ServiceType.DoesNotExist:
        return JsonResponse({'error': 'Service type not found'}, status=404)


@login_required(login_url='/login/')
def service_type_detail(request, service_type_id):
    st = get_object_or_404(ServiceType, id=service_type_id)
    form = ServiceTypeForm(instance=st)

    if request.method == 'POST':
        st.office_study = 'office_study' in request.POST
        st.defense = 'defense' in request.POST
        st.consultation = 'consultation' in request.POST
        st.check_documents = 'check_documents' in request.POST
        st.contract = 'contract' in request.POST
        st.notification = 'notification' in request.POST
        st.bill = 'bill' in request.POST

        fields_map = [
            'bill_fee', 'initial_fee', 'appeal_fee', 'enforcement_fee',
            'petition_fee', 'travel_fee', 'miscellaneous_fee', 'contract_fee',
            'notification_fee', 'court_service_fee', 'document_fee',
            'office_study_fee',
        ]
        for f in fields_map:
            setattr(st, f, Decimal(request.POST.get(f, 0) or 0))

        st.save()
        messages.success(request, '✅ اطلاعات خدمات با موفقیت ویرایش شد.')
        return redirect('movakel_detail', slug=st.movakel.slug)

    return render(request, 'movakel_module/service_type_detail.html',
                  {'service_type': st, 'form': form})


@login_required(login_url='/login/')
def service_type_edit(request, pk):
    st = get_object_or_404(ServiceType, pk=pk)
    if request.method == "POST":
        form = ServiceTypeForm(request.POST, instance=st)
        if form.is_valid():
            form.save()
            return redirect("movakel-detail", slug=st.movakel.slug)
    else:
        form = ServiceTypeForm(instance=st)

    return render(request, "movakel_module/service_type_detail.html", {
        "form": form,
        "service_type": st,
    })


@login_required(login_url='/login/')
def update_services(request, movakel_id):
    movakel = get_object_or_404(Movakel, id=movakel_id)

    if request.method == 'POST':
        st = movakel.service_type
        st.defense = 'defense' in request.POST
        st.consultation = 'consultation' in request.POST
        st.check_documents = 'check_documents' in request.POST
        st.contract = 'contract' in request.POST
        st.office_study = 'office_study' in request.POST
        st.bill = 'bill' in request.POST
        st.notification = 'notification' in request.POST

        fields_map = [
            'travel_fee', 'miscellaneous_fee', 'office_study_fee', 'contract_fee',
            'bill_fee', 'initial_fee', 'appeal_fee', 'enforcement_fee',
            'petition_fee', 'court_service_fee', 'document_fee', 'notification_fee',
        ]
        for f in fields_map:
            try:
                setattr(st, f, Decimal(request.POST.get(f, 0) or 0))
            except (ValueError, TypeError):
                pass

        st.save()

    return redirect('movakel-detail', slug=movakel.slug)


@login_required(login_url='/login/')
@user_passes_test(is_admin, login_url='/access-denied/')
def show_payment_details(request, movakel_id):
    movakel = get_object_or_404(Movakel, id=movakel_id)
    st = ServiceType.objects.filter(movakel=movakel).first()

    def to_dec(v):
        return Decimal(v) if v is not None else Decimal(0)

    selected_service_type = {
        'travel_fee': to_dec(getattr(st, 'travel_fee', None)),
        'miscellaneous_fee': to_dec(getattr(st, 'miscellaneous_fee', None)),
        'contract_fee': to_dec(getattr(st, 'contract_fee', None)),
        'bill_fee': to_dec(getattr(st, 'bill_fee', None)),
        'initial_fee': to_dec(getattr(st, 'initial_fee', None)),
        'appeal_fee': to_dec(getattr(st, 'appeal_fee', None)),
        'enforcement_fee': to_dec(getattr(st, 'enforcement_fee', None)),
        'petition_fee': to_dec(getattr(st, 'petition_fee', None)),
        'court_service_fee': to_dec(getattr(st, 'court_service_fee', None)),
        'notification_fee': to_dec(getattr(st, 'notification_fee', None)),
        'document_fee': to_dec(getattr(st, 'document_fee', None)),
    }
    selected_services = st.get_selected_services() if st else []

    payment_details = PaymentDetail.objects.filter(movakel=movakel)
    defense_documents = DefenseDocument.objects.filter(movakel=movakel)

    total_service_fee = st.total_fee if st else Decimal(0)
    total_payment_detail_amount = to_dec(
        payment_details.aggregate(total_fee=Sum('amount'))['total_fee'] or 0
    )

    meetings = RequestMeeting.objects.filter(movakel=movakel)
    meeting_fees = to_dec(
        meetings.aggregate(total_fee=Sum('meeting_fee'))['total_fee'] or 0
    )

    for m in meetings:
        m.meeting_date_shamsi = (
            jdatetime.datetime.fromgregorian(datetime=m.meeting_date).strftime('%Y/%m/%d')
            if m.meeting_date else "-"
        )

    total_amount = total_service_fee + total_payment_detail_amount + meeting_fees

    total_paid = movakel.installment_payments.filter(is_paid=True)\
        .aggregate(total=Sum('amount'))['total'] or 0
    remaining = max(Decimal(0), Decimal(movakel.total_amount) - Decimal(total_paid))

    context = {
        'movakel': movakel,
        'meetings': meetings,
        'selected_service_type': selected_service_type,
        'selected_services': selected_services,
        'total_service_fee': total_service_fee,
        'contract_fee': Decimal(0),
        'payment_details': payment_details,
        'meeting_fees': meeting_fees,
        'total_payment_detail_amount': total_payment_detail_amount,
        'defense_documents': defense_documents,
        'total_amount': total_amount,
        'total_paid': total_paid,
        'remaining_amount': remaining,
        'stamp_tax': movakel.stamp_tax(),
        'judiciary_share': movakel.judiciary_share(),
        'annual_tax': movakel.annual_tax(),
        'total_tax': movakel.total_tax(),
    }
    return render(request, 'movakel_module/payment_details.html', context)


# ============================================================
# Defence Document
# ============================================================
@login_required(login_url='/login/')
def defense_document_detail(request, pk):
    doc = get_object_or_404(DefenseDocument, pk=pk)
    return render(request, "movakel_module/defense_document_detail.html", {"doc": doc})


@login_required(login_url='/login/')
def download_document_word(request, pk):
    doc = get_object_or_404(DefenseDocument, pk=pk)
    document = Document()
    document.add_heading(doc.subject, level=1).alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT

    def add_rtl_paragraph(text, bold=False):
        p = document.add_paragraph()
        p.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(12)
        run.font.name = "Tahoma"
        rtl = OxmlElement('w:rtl')
        run._r.get_or_add_rPr().append(rtl)

    add_rtl_paragraph(f"مرحله: {doc.get_stage_display()}")
    add_rtl_paragraph(f"تاریخ ثبت: {doc.created_at.strftime('%Y/%m/%d')}")
    add_rtl_paragraph("موضوع لایحه:", bold=True)
    add_rtl_paragraph(doc.subject or "متن لایحه ثبت نشده است.")

    if doc.word_file:
        add_rtl_paragraph(f"فایل Word پیوست: {doc.word_file.name}")

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    filename = f"{doc.slug or 'document'}.docx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    document.save(response)
    return response


# ============================================================
# Convert / PDF
# ============================================================
@login_required(login_url='/login/')
def convert_meeting_to_movakel(request, meeting_id):
    meeting = get_object_or_404(RequestMeeting, id=meeting_id)

    if meeting.status == "converted":
        messages.warning(request, "این درخواست قبلاً به وکالت تبدیل شده است.")
        return redirect("meeting_list")

    movakel = Movakel.all_objects.create(
        request_meeting=meeting,
        name=meeting.full_name,
        mobile_number=meeting.mobile_number,
        file_number=meeting.phone_number,
        description="پرونده ایجاد شده از درخواست ملاقات",
        age="0",
        mechanized_number=str(uuid.uuid4())[:10],
        is_active=True,
    )
    meeting.status = "converted"
    meeting.save()

    messages.success(request, "✅ درخواست ملاقات با موفقیت به موکل تبدیل شد.")
    return redirect("movakel-detail", slug=movakel.slug)


@login_required(login_url='/login/')
def generate_pdf(request, meeting_id):
    meeting = get_object_or_404(RequestMeeting, id=meeting_id)

    safe_full_name = urllib.parse.quote(meeting.full_name.replace(" ", "_"))
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{safe_full_name}.pdf"'

    p = canvas.Canvas(response)
    p.setTitle("فرم پذیرش متقاضی")

    font_path = os.path.join(settings.BASE_DIR, "static", "fonts", "IRANSans.ttf")
    if not os.path.exists(font_path):
        return HttpResponse("⚠️ فونت ایران سنس یافت نشد.", status=500)

    pdfmetrics.registerFont(TTFont("IranSans", font_path))
    p.setFont("IranSans", 14)

    p.drawString(200, 800, "فرم پذیرش متقاضی")
    p.drawString(100, 770, f"نام و نام خانوادگی: {meeting.full_name}")
    p.drawString(100, 750, f"کد ملی: {meeting.national_id}")
    birth_year = meeting.birth_date.year if meeting.birth_date else "-"
    p.drawString(100, 730, f"سال تولد: {birth_year}")
    p.drawString(100, 710, f"نشانی: {meeting.address or '-'}")
    p.drawString(100, 690, f"کد پستی: {meeting.postal_code or '-'}")
    p.drawString(100, 670, f"تلفن ثابت: {meeting.phone_number or '-'}")
    p.drawString(100, 650, f"ایمیل: {meeting.email or '-'}")

    p.showPage()
    p.save()
    return response


# ============================================================
# Misc Views
# ============================================================
@login_required(login_url='/login/')
def calculate_damage_view(request):
    result = None
    if request.method == "POST":
        form = DamageCalculationForm(request.POST)
        if form.is_valid():
            result = calculate_damage(
                form.cleaned_data["original_amount"],
                form.cleaned_data["due_date"],
                form.cleaned_data["payment_date"],
                form.cleaned_data["inflation_rate"],
            )
    else:
        form = DamageCalculationForm()
    return render(request, "movakel_module/calculate_damage.html",
                  {"form": form, "result": result})


@login_required(login_url='/login/')
def delete_movakel(request, movakel_id):
    movakel = get_object_or_404(Movakel, id=movakel_id)
    with transaction.atomic():
        movakel.movakelpayments.all().update(is_delete=True)
        Movakel.all_objects.filter(pk=movakel.pk).update(is_delete=True)
    return redirect('movakel_list')


@login_required(login_url='/login/')
def hearing_calendar_view(request):
    today = date.today()

    today_hearings = Movakel.objects.filter(
        is_delete=False, hearing_date=today
    ).order_by('hearing_time')

    upcoming_hearings = Movakel.objects.filter(
        is_delete=False,
        hearing_date__gt=today,
        hearing_date__lte=today + timedelta(days=30),
    ).order_by('hearing_date', 'hearing_time')

    past_hearings = Movakel.objects.filter(
        is_delete=False,
        hearing_date__lt=today,
        hearing_date__gte=today - timedelta(days=30),
    ).order_by('-hearing_date')

    context = {
        'today': today,
        'today_hearings': today_hearings,
        'upcoming_hearings': upcoming_hearings,
        'past_hearings': past_hearings,
        'today_count': today_hearings.count(),
        'upcoming_count': upcoming_hearings.count(),
    }
    return render(request, 'movakel_module/hearing_calendar.html', context)


@login_required(login_url='/login/')
def advanced_search_view(request):
    query = request.GET.get('q', '').strip()
    category_id = request.GET.get('category', '')
    status = request.GET.get('status', '')
    date_from = request.GET.get('date_from', '')
    date_to = request.GET.get('date_to', '')

    movakels = Movakel.objects.filter(is_delete=False)

    if query:
        movakels = movakels.filter(
            Q(name__icontains=query) |
            Q(mechanized_number__icontains=query) |
            Q(mobile_number__icontains=query) |
            Q(file_number__icontains=query) |
            Q(description__icontains=query)
        )
    if category_id:
        movakels = movakels.filter(category_id=category_id)
    if status == 'active':
        movakels = movakels.filter(is_active=True)
    elif status == 'inactive':
        movakels = movakels.filter(is_active=False)
    if date_from:
        movakels = movakels.filter(contract_date__gte=date_from)
    if date_to:
        movakels = movakels.filter(contract_date__lte=date_to)

    context = {
        'movakels': movakels.order_by('-id'),
        'categories': MovakelCategory.objects.all(),
        'query': query,
        'selected_category': category_id,
        'selected_status': status,
        'date_from': date_from,
        'date_to': date_to,
        'result_count': movakels.count(),
    }
    return render(request, 'movakel_module/advanced_search.html', context)


@login_required(login_url='/login/')
def update_case_status(request, movakel_id):
    movakel = get_object_or_404(Movakel, id=movakel_id, is_delete=False)
    if request.method == 'POST':
        new_status = request.POST.get('case_status')
        valid = ['in_progress', 'appeal', 'execution', 'closed', 'archived']
        if new_status in valid:
            movakel.case_status = new_status
            movakel.save(update_fields=['case_status'])
            messages.success(request, f'✅ وضعیت پرونده به «{movakel.get_case_status_display()}» تغییر یافت.')
        else:
            messages.error(request, '❌ وضعیت نامعتبر است.')
    return redirect('movakel-detail', slug=movakel.slug)


@login_required(login_url='/login/')
@user_passes_test(is_admin, login_url='/access-denied/')
def backup_view(request):
    message = error = None
    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'create':
            path, result = create_backup()
            if path:
                message = f'✅ فایل پشتیبان «{result}» با موفقیت ایجاد شد.'
            else:
                error = f'❌ خطا: {result}'
        elif action == 'delete':
            filename = os.path.basename(request.POST.get('filename', ''))
            if delete_backup(filename):
                message = f'✅ فایل «{filename}» حذف شد.'
            else:
                error = '❌ فایل یافت نشد.'

    return render(request, 'movakel_module/backup.html', {
        'backups': get_all_backups(),
        'message': message,
        'error': error,
    })


@login_required(login_url='/login/')
@user_passes_test(is_admin, login_url='/access-denied/')
def download_backup(request, filename):
    backup_dir = os.path.join(settings.BASE_DIR, 'backups')

    # جلوگیری از Path Traversal: فقط نام خالص فایل مجاز است، نه مسیر
    safe_filename = os.path.basename(filename)
    file_path = os.path.join(backup_dir, safe_filename)

    # اطمینان مضاعف از اینکه مسیر نهایی هنوز داخل پوشه backups است
    if (
        safe_filename.endswith('.sqlite3')
        and os.path.commonpath([os.path.abspath(file_path), os.path.abspath(backup_dir)]) == os.path.abspath(backup_dir)
        and os.path.exists(file_path)
    ):
        return FileResponse(open(file_path, 'rb'), as_attachment=True, filename=safe_filename)

    messages.error(request, '❌ فایل یافت نشد.')
    return redirect('backup_view')


@login_required
def print_contract(request, movakel_id):
    movakel = get_object_or_404(Movakel, id=movakel_id, is_delete=False)

    font_path = os.path.join(settings.BASE_DIR, 'static', 'fonts', 'IRANSans.ttf')
    pdfmetrics.registerFont(TTFont('IRANSans', font_path))

    response = HttpResponse(content_type='application/pdf')
    safe_name = urllib.parse.quote(movakel.name.replace(' ', '_'))
    response['Content-Disposition'] = f'attachment; filename="contract_{safe_name}.pdf"'

    p = canvas.Canvas(response, pagesize=A4)
    width, height = A4
    p.setFont('IRANSans', 14)

    def draw_rtl(text, y, size=12):
        p.setFont('IRANSans', size)
        text = str(text) if text else "-"
        if BIDI_AVAILABLE:
            text = get_display(arabic_reshaper.reshape(text))
        p.drawRightString(width - 50, y, text)

    draw_rtl('قرارداد وکالت', height - 60, size=18)
    p.line(50, height - 75, width - 50, height - 75)

    draw_rtl('مشخصات موکل', height - 110, size=14)
    draw_rtl(f'نام و نام خانوادگی: {movakel.name}', height - 140)
    draw_rtl(f'شماره تماس: {movakel.mobile_number}', height - 165)
    draw_rtl(f'شماره پرونده مکانیزه: {movakel.mechanized_number or "-"}', height - 190)
    draw_rtl(f'تاریخ وکالت: {movakel.contract_date or "-"}', height - 215)
    draw_rtl(f'دسته‌بندی: {movakel.category or "-"}', height - 240)

    p.line(50, height - 255, width - 50, height - 255)
    draw_rtl('مشخصات مالی', height - 285, size=14)
    draw_rtl(f'مبلغ کل حق‌الوکاله: {movakel.total_amount:,} تومان', height - 315)
    draw_rtl(f'نحوه پرداخت: {movakel.get_payment_method_display() if movakel.payment_method else "-"}', height - 340)

    p.line(50, height - 355, width - 50, height - 355)
    draw_rtl('موضوع وکالت', height - 385, size=14)
    draw_rtl(movakel.description[:100] if movakel.description else '-', height - 415, size=11)

    p.line(50, height - 430, width - 50, height - 430)
    draw_rtl('امضای موکل', height - 530, size=11)
    draw_rtl('امضای وکیل', height - 530, size=11)
    p.line(50, height - 510, 200, height - 510)
    p.line(width - 200, height - 510, width - 50, height - 510)

    draw_rtl('این قرارداد با استفاده از سامانه وکیل‌یار تنظیم شده است.', 40, size=9)

    p.showPage()
    p.save()
    return response


# ============================================================
# Site Components
# ============================================================
def site_header_component(request):
    setting = SiteSetting.objects.filter(is_main_setting=True).first()
    return render(request, 'shared/site_header_component.html', {'site_setting': setting})


def site_footer_component(request):
    setting = SiteSetting.objects.filter(is_main_setting=True).first()
    try:
        from site_module.models import FooterLinkBox
        boxes = FooterLinkBox.objects.all()
    except (ImportError, Exception):
        boxes = []
    return render(request, 'shared/site_footer_component.html', {
        'site_setting': setting,
        'footer_link_boxes': boxes,
    })
    
    
    
@login_required
def total_costs_view(request):
    """
    گزارش مالی کلی از تمام پرونده‌ها
    """
    # ─── مجموع مبلغ توافق شده ───
    total_total_amount = Movakel.objects.filter(
        is_delete=False
    ).aggregate(total=Sum('total_amount'))['total'] or 0
    
    # ─── مجموع پرداخت‌های انجام شده ───
    total_received_amount = MovakelPayment.objects.filter(
        is_delete=False
    ).aggregate(total=Sum('amount'))['total'] or 0
    
    # ─── مجموع باقی‌مانده ───
    total_remaining_amount = total_total_amount - total_received_amount
    
    # ─── مجموع هزینه‌های خدمات ───
    # ServiceType فیلد total_fee نداره، پس دستی جمع می‌کنیم
    total_service_cost = 0
    for st in ServiceType.objects.all():
        try:
            total_service_cost += st.total_fee or 0
        except Exception:
            pass
    
    # ─── مجموع پرداختی‌ها ───
    total_payments = total_received_amount
    
    # ─── جزئیات پرداختی‌ها بر اساس نوع هزینه ───
    payments_by_category = MovakelPayment.objects.filter(
        is_delete=False
    ).values('payment_for').annotate(
        total=Sum('amount')
    ).order_by('-total')
    
    context = {
        'total_total_amount': total_total_amount,
        'total_received_amount': total_received_amount,
        'total_remaining_amount': total_remaining_amount,
        'total_service_cost': total_service_cost,
        'total_payments': total_payments,
        'payments_by_category': payments_by_category,
    }
    
    return render(request, 'movakel_module/total_costs.html', context)
    """
    گزارش مالی کلی از تمام پرونده‌ها
    """
    # ─── مجموع مبلغ توافق شده ───
    total_total_amount = Movakel.objects.filter(
        is_delete=False
    ).aggregate(total=Sum('total_amount'))['total'] or 0
    
    # ─── مجموع پرداخت‌های انجام شده ───
    total_received_amount = MovakelPayment.objects.filter(
        is_delete=False
    ).aggregate(total=Sum('amount'))['total'] or 0
    
    # ─── مجموع باقی‌مانده ───
    total_remaining_amount = total_total_amount - total_received_amount
    
    # ─── مجموع هزینه‌های خدمات ───
    total_service_cost = ServiceType.objects.aggregate(
        total=Sum('total_fee')
    )['total'] or 0
    
    # ─── مجموع پرداختی‌ها ───
    total_payments = total_received_amount
    
    # ─── جزئیات پرداختی‌ها بر اساس نوع هزینه ───
    payments_by_category = MovakelPayment.objects.filter(
        is_delete=False
    ).values('payment_for').annotate(
        total=Sum('amount')
    ).order_by('-total')
    
    context = {
        'total_total_amount': total_total_amount,
        'total_received_amount': total_received_amount,
        'total_remaining_amount': total_remaining_amount,
        'total_service_cost': total_service_cost,
        'total_payments': total_payments,
        'payments_by_category': payments_by_category,
    }
    
    return render(request, 'movakel_module/total_costs.html', context) 



# ============================================================
# PDF Generation Views
# ============================================================

@login_required(login_url='/login/')
def pdf_contract(request, movakel_id):
    """تولید PDF قرارداد وکالت"""
    movakel = get_object_or_404(Movakel, id=movakel_id, is_delete=False)

    context = get_common_pdf_context()
    context.update({
        'doc_type': 'contract',
        'movakel': movakel,
    })

    # خدمات انتخاب‌شده (اگه ServiceType داره)
    try:
        context['selected_services'] = movakel.service_type.get_selected_services()
    except ServiceType.DoesNotExist:
        context['selected_services'] = []

    pdf_content = render_to_pdf('movakel_module/pdf.html', context, request)
    if pdf_content is None:
        messages.error(request, '❌ خطا در تولید PDF. لطفاً دوباره تلاش کنید.')
        return redirect('movakel-detail', slug=movakel.slug)

    filename = safe_pdf_filename('contract', movakel.name)
    response = HttpResponse(pdf_content.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{filename}"'
    return response


@login_required(login_url='/login/')
def pdf_report(request, movakel_id):
    """تولید PDF گزارش مالی پرونده"""
    movakel = get_object_or_404(Movakel, id=movakel_id, is_delete=False)

    # پرداخت‌ها
    payments = movakel.movakelpayments.filter(is_delete=False).order_by('-payment_date')
    total_paid = payments.aggregate(total=Sum('amount'))['total'] or 0
    remaining_amount = max(0, movakel.total_amount - total_paid)

    # اقساط
    installments = movakel.installment_payments.all().order_by('due_date')

    # سرویس تایپ
    try:
        service_type = movakel.service_type
        total_service_fee = service_type.total_fee
    except ServiceType.DoesNotExist:
        service_type = None
        total_service_fee = 0

    # مالیات‌ها
    stamp_tax = movakel.stamp_tax()
    judiciary_share = movakel.judiciary_share()
    annual_tax = movakel.annual_tax()
    total_tax = movakel.total_tax()

    context = get_common_pdf_context()
    context.update({
        'doc_type': 'report',
        'movakel': movakel,
        'payments': payments,
        'total_paid': total_paid,
        'remaining_amount': remaining_amount,
        'installments': installments,
        'service_type': service_type,
        'total_service_fee': total_service_fee,
        'stamp_tax': stamp_tax,
        'judiciary_share': judiciary_share,
        'annual_tax': annual_tax,
        'total_tax': total_tax,
    })

    pdf_content = render_to_pdf('movakel_module/pdf.html', context, request)
    if pdf_content is None:
        messages.error(request, '❌ خطا در تولید PDF.')
        return redirect('movakel-detail', slug=movakel.slug)

    filename = safe_pdf_filename('report', movakel.name)
    response = HttpResponse(pdf_content.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{filename}"'
    return response


@login_required(login_url='/login/')
def pdf_total_report(request):
    """تولید PDF گزارش مالی کلی دفتر"""
    from .models import MovakelPayment

    # مجموع مبلغ توافق‌شده
    total_total_amount = Movakel.objects.filter(
        is_delete=False
    ).aggregate(total=Sum('total_amount'))['total'] or 0

    # مجموع پرداخت‌ها
    total_received_amount = MovakelPayment.objects.filter(
        is_delete=False
    ).aggregate(total=Sum('amount'))['total'] or 0

    total_remaining_amount = max(0, total_total_amount - total_received_amount)

    # مجموع هزینه‌های خدمات
    total_service_cost = 0
    for st in ServiceType.objects.all():
        try:
            total_service_cost += st.total_fee or 0
        except Exception:
            pass

    # پرداخت‌ها بر اساس نوع هزینه
    payments_by_category = MovakelPayment.objects.filter(
        is_delete=False
    ).values('payment_for').annotate(
        total=Sum('amount')
    ).order_by('-total')

    # تبدیل کد payment_for به عنوان فارسی
    payment_for_map = dict(PAYMENT_FOR_CHOICES)
    payments_by_category_display = []
    for item in payments_by_category:
        payments_by_category_display.append({
            'payment_for': payment_for_map.get(item['payment_for'], item['payment_for']),
            'total': item['total'],
        })

    context = get_common_pdf_context()
    context.update({
        'doc_type': 'total_report',
        'total_total_amount': total_total_amount,
        'total_received_amount': total_received_amount,
        'total_remaining_amount': total_remaining_amount,
        'total_service_cost': total_service_cost,
        'payments_by_category': payments_by_category_display,
    })

    pdf_content = render_to_pdf('movakel_module/pdf.html', context, request)
    if pdf_content is None:
        messages.error(request, '❌ خطا در تولید PDF.')
        return redirect('total_costs')

    filename = safe_pdf_filename('total_report', timezone.now().strftime('%Y%m%d'))
    response = HttpResponse(pdf_content.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{filename}"'
    return response


@login_required(login_url='/login/')
def pdf_meeting_form(request, meeting_id):
    """تولید PDF فرم پذیرش متقاضی"""
    meeting = get_object_or_404(RequestMeeting, id=meeting_id)

    context = get_common_pdf_context()
    context.update({
        'doc_type': 'meeting_form',
        'meeting': meeting,
    })

    pdf_content = render_to_pdf('movakel_module/pdf.html', context, request)
    if pdf_content is None:
        messages.error(request, '❌ خطا در تولید PDF.')
        return redirect('meeting_list')

    filename = safe_pdf_filename('meeting', meeting.full_name)
    response = HttpResponse(pdf_content.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{filename}"'
    return response


